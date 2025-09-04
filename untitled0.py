import re
import io
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image
from openpyxl import load_workbook

# =========================
# Configuración de la app
# =========================
st.set_page_config(page_title="SIAF Dashboard - Peru Compras", layout="wide")
st.title("SIAF Dashboard - Peru Compras")

st.markdown(
    "Carga el **Excel SIAF** y obtén resúmenes de **PIA, PIM, Certificado, Comprometido, Devengado, Saldo PIM y % Avance**. "
    "Lee por defecto **A:EC** (base hasta CH + pasos CI–EC). "
    "Construye el **clasificador concatenado** y lo **normaliza para que siempre comience con `2.`**; además agrega una **descripción jerárquica**. "
    "Incluye filtros, pivotes, serie mensual y descarga a Excel."
)

# =========================
# Sidebar / parámetros
# =========================
with st.sidebar:
    st.header("Parámetros de lectura")
    uploaded = st.file_uploader("Archivo SIAF (.xlsx)", type=["xlsx"])
    usecols = st.text_input("Rango de columnas (Excel)", "A:EC", help="Recomendado para incluir CI–EC.")
    sheet_name = st.text_input("Nombre de hoja (opcional)", "", help="Déjalo vacío para autodetección.")
    header_row_excel = st.number_input("Fila de encabezados (Excel, 1=primera)", min_value=1, value=4)
    detect_header = st.checkbox("Autodetectar encabezado", value=True)
    st.markdown("---")
    st.header("Reglas CI–EC")
    current_month = st.number_input("Mes actual (1-12)", min_value=1, max_value=12, value=9)
    riesgo_umbral = st.number_input("Umbral de avance mínimo (%)", min_value=0, max_value=100, value=60)
    st.caption("Se marca riesgo_devolucion si Avance% < Umbral.")

# =========================
# Utilitarios de carga
# =========================
def autodetect_sheet_and_header(xls, excel_bytes, usecols, user_sheet, header_guess):
    """
    Busca la hoja y la fila que luce como encabezado (contenga 'ano_eje', 'mto_', 'pim', etc.).
    Retorna (sheet_name, header_row_index_pandas).
    """
    candidate_sheets = [user_sheet] if user_sheet else xls.sheet_names
    for s in candidate_sheets:
        try:
            tmp = pd.read_excel(excel_bytes, sheet_name=s, header=None, usecols=usecols, nrows=12)
        except Exception:
            continue
        for r in range(min(8, len(tmp))):
            row_vals = tmp.iloc[r].astype(str).str.lower().tolist()
            hits = sum(int(any(k in v for k in ["ano_eje", "pim", "pia", "mto_", "devenga", "girado"])) for v in row_vals)
            if hits >= 2:
                return s, r
    # Fallback: primera hoja y fila indicada por el usuario - 1 (a índice 0)
    return xls.sheet_names[0], header_guess - 1

def load_data(excel_bytes, usecols, sheet_name, header_row_excel, autodetect=True):
    xls = pd.ExcelFile(excel_bytes)
    if autodetect:
        s, hdr = autodetect_sheet_and_header(xls, excel_bytes, usecols, sheet_name, header_row_excel)
        df = pd.read_excel(excel_bytes, sheet_name=s, header=hdr, usecols=usecols)
    else:
        hdr = header_row_excel - 1
        s = sheet_name if sheet_name else xls.sheet_names[0]
        df = pd.read_excel(excel_bytes, sheet_name=s, header=hdr, usecols=usecols)

    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip().lower() for c in df.columns]
    return df, s

# =========================
# Cálculos CI–EC
# =========================
def find_monthly_columns(df, prefix):
    return [f"{prefix}{i:02d}" for i in range(1, 13) if f"{prefix}{i:02d}" in df.columns]

def ensure_ci_ec_steps(df, month, umbral):
    """
    Crea/asegura columnas claves si no existen:
    - devengado (suma mto_devenga_01..12)
    - devengado_mes (columna del mes seleccionado)
    - saldo_pim (pim - devengado)
    - avance_% (devengado/pim)
    - riesgo_devolucion (avance_% < umbral)
    - area (vacía si no existe)
    """
    df = df.copy()
    dev_cols = find_monthly_columns(df, "mto_devenga_")

    if "devengado" not in df.columns:
        df["devengado"] = df[dev_cols].sum(axis=1) if dev_cols else 0.0

    col_mes = f"mto_devenga_{int(month):02d}"
    if "devengado_mes" not in df.columns:
        df["devengado_mes"] = df[col_mes] if col_mes in df.columns else 0.0

    if "saldo_pim" not in df.columns:
        df["saldo_pim"] = np.where(df.get("mto_pim", 0) > 0, df["mto_pim"] - df["devengado"], 0.0)

    if "avance_%" not in df.columns:
        df["avance_%"] = np.where(df.get("mto_pim", 0) > 0, df["devengado"] / df["mto_pim"] * 100.0, 0.0)

    if "riesgo_devolucion" not in df.columns:
        df["riesgo_devolucion"] = df["avance_%"] < float(umbral)

    if "area" not in df.columns:
        df["area"] = ""

    return df

# =========================
# Clasificador concatenado
# =========================
_code_re = re.compile(r"^\s*(\d+(?:\.\d+)*)")

def extract_code(text):
    """Extrae el prefijo numérico (con puntos) de un texto tipo '2.1.1 Bienes y servicios'."""
    if pd.isna(text):
        return ""
    s = str(text).strip()
    m = _code_re.match(s)
    return m.group(1) if m else ""

def last_segment(code):
    return code.split(".")[-1] if code else ""

def concat_hierarchy(gen, sub, subdet, esp, espdet):
    """
    Concatena jerárquicamente evitando duplicados:
    generica.subgenerica.subgenerica_det.especifica.especifica_det
    """
    parts = []
    if gen:
        parts.append(gen)
    for child in [sub, subdet, esp, espdet]:
        if not child:
            continue
        # Si el hijo ya trae el prefijo, lo conservamos
        if parts and (child.startswith(parts[-1] + ".") or child.startswith(parts[0] + ".")):
            parts.append(child)
        else:
            # Caso contrario, agregamos solo el último segmento al prefijo anterior
            if parts:
                parts.append(parts[-1] + "." + last_segment(child))
            else:
                parts.append(child)
    return parts[-1] if parts else ""

def normalize_clasificador(code):
    """
    Regla: todo clasificador debe comenzar con '2.'.
    - Si está vacío => '2.'
    - Si no inicia con '2.' => anteponer '2.'
    """
    if not code:
        return "2."
    return code if code.startswith("2.") else "2." + code

def desc_only(text):
    """Devuelve solo la descripción (lo que va después del primer punto)."""
    if pd.isna(text):
        return ""
    s = str(text)
    return s.split(".", 1)[1].strip() if "." in s else s

def build_classifier_columns(df):
    """
    Crea columnas:
    - gen_cod, sub_cod, subdet_cod, esp_cod, espdet_cod (códigos numéricos)
    - clasificador_cod (concatenado y normalizado con 2.)
    - generica_desc, subgenerica_desc, subgenerica_det_desc, especifica_desc, especifica_det_desc
    - clasificador_desc (descripción jerárquica)
    """
    df = df.copy()
    gen = df.get("generica", "")
    sub = df.get("subgenerica", "")
    subdet = df.get("subgenerica_det", "")
    esp = df.get("especifica", "")
    espdet = df.get("especifica_det", "")

    df["gen_cod"] = gen.map(extract_code) if "generica" in df.columns else ""
    df["sub_cod"] = sub.map(extract_code) if "subgenerica" in df.columns else ""
    df["subdet_cod"] = subdet.map(extract_code) if "subgenerica_det" in df.columns else ""
    df["esp_cod"] = esp.map(extract_code) if "especifica" in df.columns else ""
    df["espdet_cod"] = espdet.map(extract_code) if "especifica_det" in df.columns else ""

    df["clasificador_cod"] = [
        normalize_clasificador(concat_hierarchy(g, s, sd, e, ed))
        for g, s, sd, e, ed in zip(
            df["gen_cod"], df["sub_cod"], df["subdet_cod"], df["esp_cod"], df["espdet_cod"]
        )
    ]

    # Descripciones sin código
    df["generica_desc"] = gen.map(desc_only) if "generica" in df.columns else ""
    df["subgenerica_desc"] = sub.map(desc_only) if "subgenerica" in df.columns else ""
    df["subgenerica_det_desc"] = subdet.map(desc_only) if "subgenerica_det" in df.columns else ""
    df["especifica_desc"] = esp.map(desc_only) if "especifica" in df.columns else ""
    df["especifica_det_desc"] = espdet.map(desc_only) if "especifica_det" in df.columns else ""

    df["clasificador_desc"] = (
        df["generica_desc"].fillna("")
        + " > " + df["subgenerica_desc"].fillna("")
        + " > " + df["subgenerica_det_desc"].fillna("")
        + " > " + df["especifica_desc"].fillna("")
        + " > " + df["especifica_det_desc"].fillna("")
    ).str.strip(" >")

    return df

# =========================
# Pivote / resumen por grupo
# =========================
def pivot_exec(df, group_col, dev_cols):
    cols = []
    if "mto_pia" in df.columns:
        cols.append("mto_pia")
    if "mto_pim" in df.columns:
        cols.append("mto_pim")
    if "mto_certificado" in df.columns:
        cols.append("mto_certificado")
    if "mto_compro_anual" in df.columns:
        cols.append("mto_compro_anual")
    if dev_cols:
        cols.append("devengado")

    # Si no existía 'devengado' pero hay columnas mensuales, lo armamos en copia
    if "devengado" not in df.columns and dev_cols:
        df = df.copy()
        df["devengado"] = df[dev_cols].sum(axis=1)

    g = df.groupby(group_col, dropna=False)[cols].sum().reset_index()

    if "mto_pim" in g.columns and "devengado" in g.columns:
        g["saldo_pim"] = g["mto_pim"] - g["devengado"]
        g["avance_%"] = np.where(g["mto_pim"] > 0, g["devengado"] / g["mto_pim"] * 100.0, 0.0)

    return g

# =========================
# Styling Function
# =========================
def highlight_risk(row, umbral):
    """Highlights rows where 'avance_%' is below the specified umbral."""
    if 'avance_%' in row and isinstance(row['avance_%'], (int, float)) and row['avance_%'] < umbral:
        return ['background-color: yellow'] * len(row)
    return [''] * len(row)


# =========================
# Carga del archivo
# =========================
if uploaded is None:
    st.info("Sube tu archivo Excel SIAF para empezar. Usa A:EC para incluir pasos CI–EC.")
    st.stop()

try:
    df, used_sheet = load_data(uploaded, usecols, sheet_name.strip() or None, int(header_row_excel), autodetect=detect_header)
except Exception as e:
    st.error(f"No se pudo leer el archivo: {e}")
    st.stop()

st.success(f"Leída la hoja '{used_sheet}' con {df.shape[0]} filas y {df.shape[1]} columnas.")

# =========================
# Filtros
# =========================
st.subheader("Filtros")
filter_cols = [c for c in df.columns if any(k in c for k in [
    "unidad_ejecutora","fuente_financ","generica","subgenerica","subgenerica_det",
    "especifica","especifica_det","funcion","division_fn","grupo_fn","programa_pptal",
    "producto_proyecto","activ_obra_accinv","meta","sec_func",
    "departamento_meta","provincia_meta","distrito_meta","area"
])]

cols_f = st.columns(3)
selected_filters = {}
for i, c in enumerate(filter_cols):
    with cols_f[i % 3]:
        vals = sorted([str(x) for x in df[c].dropna().unique().tolist()])
        if len(vals) > 1:
            pick = st.multiselect(c, options=vals, default=[])
            if pick:
                selected_filters[c] = set(pick)

mask = pd.Series(True, index=df.index)
for c, allowed in selected_filters.items():
    mask &= df[c].astype(str).isin(allowed)
df_f = df[mask].copy()

# =========================
# Aplicar CI–EC + Clasificador
# =========================
df_proc = ensure_ci_ec_steps(df_f, current_month, riesgo_umbral)
df_proc = build_classifier_columns(df_proc)

# =========================
# Resumen ejecutivo
# =========================
st.subheader("Resumen ejecutivo (totales)")
dev_cols = [c for c in df_proc.columns if c.startswith("mto_devenga_")]

tot_pia = float(df_proc.get("mto_pia", 0).sum())
tot_pim = float(df_proc.get("mto_pim", 0).sum())
tot_dev = float(df_proc.get("devengado", 0).sum())
tot_cert = float(df_proc.get("mto_certificado", 0).sum()) if "mto_certificado" in df_proc.columns else 0.0
tot_comp = float(df_proc.get("mto_compro_anual", 0).sum()) if "mto_compro_anual" in df_proc.columns else 0.0
saldo_pim = tot_pim - tot_dev if tot_pim else 0.0
avance = (tot_dev / tot_pim * 100.0) if tot_pim else 0.0

k1, k2, k3, k4, k5, k6, k7 = st.columns(7)
k1.metric("PIA", f"S/ {tot_pia:,.2f}")
k2.metric("PIM", f"S/ {tot_pim:,.2f}")
k3.metric("Certificado", f"S/ {tot_cert:,.2f}")
k4.metric("Comprometido", f"S/ {tot_comp:,.2f}")
k5.metric("Devengado (YTD)", f"S/ {tot_dev:,.2f}")
k6.metric("Saldo PIM", f"S/ {saldo_pim:,.2f}")
k7.metric("Avance", f"{avance:.2f}%")

# =========================
# Vistas por agrupación con highlighting
# =========================
st.subheader("Vistas por agrupación")
group_options = [c for c in df_proc.columns if c in [
    "clasificador_cod","unidad_ejecutora","fuente_financ","generica","subgenerica","subgenerica_det",
    "especifica","especifica_det","funcion","division_fn","grupo_fn","programa_pptal",
    "producto_proyecto","activ_obra_accinv","meta","sec_func","area"
]]
default_idx = group_options.index("clasificador_cod") if "clasificador_cod" in group_options else 0
group_col = st.selectbox("Agrupar por", options=group_options, index=default_idx)

pivot = pivot_exec(df_proc, group_col, dev_cols)
styled_pivot = pivot.style.apply(highlight_risk, umbral=riesgo_umbral, axis=1)
st.dataframe(styled_pivot, width='stretch') # Updated width

# =========================
# Procesos CI–EC (detalle) con highlighting
# =========================
st.subheader("Procesos CI–EC (monto vinculado a su cadena)")
ci_cols = [
    "clasificador_cod", "clasificador_desc",
    "generica","subgenerica","subgenerica_det","especifica","especifica_det",
    "mto_pia","mto_pim","mto_certificado","mto_compro_anual",
    "devengado_mes","devengado","saldo_pim","avance_%","riesgo_devolucion"
]
ci_cols = [c for c in ci_cols if c in df_proc.columns]
ci_cols_df_head = df_proc[ci_cols].head(300).copy()
styled_ci_cols_df_head = ci_cols_df_head.style.apply(highlight_risk, umbral=riesgo_umbral, axis=1)
st.dataframe(styled_ci_cols_df_head, width='stretch') # Updated width

# =========================
# Consolidado por clasificador con highlighting
# =========================
agg_cols = [c for c in ["mto_pia","mto_pim","mto_certificado","mto_compro_anual","devengado_mes","devengado","saldo_pim"] if c in df_proc.columns]
consolidado = df_proc.groupby(
    ["clasificador_cod","clasificador_desc","generica","subgenerica","subgenerica_det","especifica","especifica_det"],
    dropna=False
)[agg_cols].sum().reset_index()

if "mto_pim" in consolidado.columns and "devengado" in consolidado.columns:
    consolidado["avance_%"] = np.where(consolidado["mto_pim"] > 0, consolidado["devengado"]/consolidado["mto_pim"]*100.0, 0.0)

st.markdown("**Consolidado por clasificador**")
consolidado_head = consolidado.head(500)
styled_consolidado_head = consolidado_head.style.apply(highlight_risk, umbral=riesgo_umbral, axis=1)
st.dataframe(styled_consolidado_head, width='stretch') # Updated width


# =========================
# Serie mensual de devengado (dinámica e interactiva)
# =========================
if dev_cols:
    st.subheader("Devengado mensual (por filtro actual)")

    # Add filter for the selected group_col
    all_categories = ["Todos"] + sorted(df_proc[group_col].dropna().unique().tolist())
    selected_category = st.selectbox(f"Filtrar por {group_col}", options=all_categories, index=0)

    if selected_category != "Todos":
        df_dev_filtered = df_proc[df_proc[group_col] == selected_category].copy()
    else:
        df_dev_filtered = df_proc.copy()

    month_map = {f"mto_devenga_{i:02d}": i for i in range(1, 13)}
    dev_series = df_dev_filtered[dev_cols].sum().reset_index()
    dev_series.columns = ["col", "monto"]
    dev_series["mes"] = dev_series["col"].map(month_map)
    dev_series = dev_series.sort_values("mes")
    st.dataframe(dev_series[["mes", "monto"]], width='stretch') # Updated width

    # Gráfico de barras simple con matplotlib (tamaño ajustado)
    fig, ax = plt.subplots(figsize=(12, 7)) # Adjusted figure size
    ax.bar(dev_series["mes"].astype(int), dev_series["monto"].values)
    ax.set_xlabel("Mes")
    ax.set_ylabel("Devengado (S/)")
    ax.set_title(f"Devengado mensual (acumulado para '{selected_category}' en '{group_col}')" if selected_category != "Todos" else "Devengado mensual (acumulado por filtro actual)")
    st.pyplot(fig)

# =========================
# Descarga a Excel (Modificada)
# =========================
def to_excel_download_modified(df_proc, pivot, consolidado, dev_series, fig, group_col, riesgo_umbral):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        # Add the plot image to a sheet
        workbook = writer.book
        if "Grafico Mensual" not in workbook.sheetnames:
             workbook.create_sheet("Grafico Mensual")
        ws_plot = workbook["Grafico Mensual"]

        img_buffer = io.BytesIO()
        fig.savefig(img_buffer, format='png')
        img_buffer.seek(0)
        plt.close(fig) # Close the figure to free memory

        img = Image(img_buffer)
        ws_plot.add_image(img, 'A1')

        # Create and format the executive summary table based on pivot
        executive_summary_cols = [col for col in ['clasificador_cod', 'clasificador_desc', group_col, 'mto_pia', 'mto_pim', 'mto_certificado', 'mto_compro_anual', 'devengado', 'saldo_pim', 'avance_%'] if col in pivot.columns]
        executive_summary_df = pivot[executive_summary_cols].copy()

        # Apply highlighting to the executive summary table data before writing
        styled_executive_summary_df = executive_summary_df.style.apply(highlight_risk, umbral=riesgo_umbral, axis=1)

        # Write the styled executive summary DataFrame to a new sheet
        # Note: Writing a styled DataFrame directly to Excel might lose some formatting.
        # A common approach is to write the raw data and apply formatting using openpyxl
        # after writing. For simplicity here, we'll write the raw data and mention this limitation.
        executive_summary_df.to_excel(writer, index=False, sheet_name="Resumen Ejecutivo")

        # Optional: Add a note in the Excel about the styling not being preserved automatically
        # ws_summary = workbook["Resumen Ejecutivo"]
        # ws_summary['A1'] = "Note: Conditional highlighting from the app is not preserved in this raw data export."


    output.seek(0)
    return output

# Download button using the modified function
if uploaded is not None and 'df_proc' in locals() and 'pivot' in locals() and 'dev_series' in locals() and 'fig' in locals():
    buf_modified = to_excel_download_modified(
        df_proc=df_proc,
        pivot=pivot,
        consolidado=consolidado, # consolidado is included for potential future use, but not written in this modified function
        dev_series=dev_series,
        fig=fig,
        group_col=group_col,
        riesgo_umbral=riesgo_umbral
    )
    st.download_button(
        "Descargar Excel (Gráfico y Resumen Ejecutivo)",
        data=buf_modified,
        file_name="siaf_analisis_ejecutivo.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )